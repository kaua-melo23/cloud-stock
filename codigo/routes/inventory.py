import traceback
import os
import uuid
import re
import pytz
from datetime import datetime
from flask import Blueprint, request, jsonify
from flask_login import current_user, login_required
from openpyxl import load_workbook
from database import db, audit_log
from auth import is_analyst_or_admin, is_admin
from config import gsheets, settings

bp = Blueprint('inventory_sys', __name__)

UPLOAD_FOLDER = os.path.join(os.getcwd(), 'uploads_temp')
if not os.path.exists(UPLOAD_FOLDER): os.makedirs(UPLOAD_FOLDER)

def get_now_br(): 
    return datetime.now(pytz.timezone('America/Recife')).strftime('%Y-%m-%d %H:%M:%S')

def parse_date(s):
    if not str(s).strip() or str(s).lower() in ['none','','nan','-']: return datetime.min
    for fmt in ['%Y-%m-%d %H:%M:%S','%d/%m/%Y %H:%M:%S','%Y-%m-%d','%d/%m/%Y']:
        try: return datetime.strptime(str(s).split('.')[0], fmt)
        except: continue
    return datetime.min

@bp.route('/api/inventory', methods=['GET'])
@login_required
def list_inventory():
    # Retorna itens que não estão marcados para exclusão
    items = db.get_connection('inventory').execute("SELECT * FROM inventory WHERE status!='EXCLUIR' ORDER BY id DESC").fetchall()
    return jsonify([dict(i) for i in items])

@bp.route('/api/inventory/search-available', methods=['GET'])
@login_required
def search_available_item():
    q = request.args.get('q', '').strip()
    if not q: return jsonify({'found': False, 'message': 'Digite algo.'}), 400
    
    conn = db.get_connection('inventory')
    # Busca Case-Insensitive (LIKE)
    item = conn.execute('SELECT * FROM inventory WHERE (serial_number LIKE ? OR tag LIKE ? OR tombamento LIKE ?) AND status="Available"', (q,q,q)).fetchone()
    
    if item: return jsonify({'found': True, 'item': dict(item)})
    
    check = conn.execute('SELECT status FROM inventory WHERE serial_number LIKE ? OR tag LIKE ? OR tombamento LIKE ?', (q,q,q)).fetchone()
    if check: return jsonify({'found': False, 'message': f'Item encontrado, mas está "{check["status"]}".'}), 400
    
    return jsonify({'found': False, 'message': 'Item não encontrado no estoque.'}), 404

@bp.route('/api/inventory/manual', methods=['POST'])
@is_admin
def add_manual():
    d = request.get_json()
    conn = db.get_connection('inventory')
    try:
        conn.execute('INSERT INTO inventory (serial_number,type,model,tag,tombamento,condition,status,glpi_status,standardized,last_updated) VALUES (?,?,?,?,?,?,?,?,?,?)', 
                     (d.get('serial_number'),d.get('type'),d.get('model'),d.get('tag'),d.get('tombamento'),'Novo',d.get('status','Available'),d.get('glpi_status'),d.get('standardized'),get_now_br()))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e: return jsonify({'message': str(e)}), 500

@bp.route('/api/inventory/<int:id>', methods=['PUT'])
@is_admin
def update_item(id):
    d = request.get_json()
    conn = db.get_connection('inventory')
    try:
        conn.execute('''UPDATE inventory SET type=?, model=?, serial_number=?, tombamento=?, status=?, glpi_status=?, standardized=?, last_updated=? WHERE id=?''',
            (d.get('type'), d.get('model'), d.get('serial_number'), d.get('tombamento'), d.get('status'), d.get('glpi_status'), d.get('standardized'), get_now_br(), id))
        conn.commit()
        return jsonify({'success': True})
    except Exception as e: return jsonify({'message': str(e)}), 500

@bp.route('/api/inventory/<int:id>', methods=['DELETE'])
@is_admin
def delete_inv(id):
    db.get_connection('inventory').execute("UPDATE inventory SET status='EXCLUIR', last_updated=? WHERE id=?", (get_now_br(), id)).connection.commit()
    return jsonify({'success': True})

@bp.route('/api/inventory/batch-delete', methods=['POST'])
@is_admin
def batch_delete():
    ids = request.get_json().get('ids', [])
    if ids:
        conn = db.get_connection('inventory')
        placeholders = ','.join(['?']*len(ids))
        conn.execute(f"UPDATE inventory SET status='EXCLUIR', last_updated=? WHERE id IN ({placeholders})", [get_now_br()] + [str(x) for x in ids])
        conn.commit()
    return jsonify({'success': True})

@bp.route('/api/inventory/sync', methods=['POST'])
@is_analyst_or_admin
def sync_inventory_google():
    conn = db.get_connection('inventory')
    try:
        db_rows = conn.execute("SELECT * FROM inventory").fetchall()
        db_map = {row['serial_number']: dict(row) for row in db_rows}
        sheet_values = gsheets.read_sheet(f"{settings.INVENTORY_RANGE}!A2:L")
        sheet_map = {}
        if sheet_values:
            for row in sheet_values:
                while len(row)<12: row.append("")
                sn = str(row[0]).strip()
                if sn: sheet_map[sn] = {'serial_number':sn,'type':row[1],'model':row[2],'tag':row[3],'tombamento':row[4],'condition':row[5],'status':row[6],'assigned_request_id':row[7],'glpi_status':row[10],'last_updated':row[11]}
        
        final, now = {}, get_now_br()
        ins, upd, dele = 0, 0, 0
        for sn in set(db_map)|set(sheet_map):
            d, s = db_map.get(sn), sheet_map.get(sn)
            w = s if d and s and parse_date(s['last_updated']) > parse_date(d['last_updated']) else (d or s)
            if str(w['status']).upper().strip() == 'EXCLUIR':
                if d: 
                    conn.execute("DELETE FROM inventory WHERE serial_number=?",(sn,))
                    dele+=1
                continue
            final[sn] = w
            if w == s: 
                exists = conn.execute("SELECT 1 FROM inventory WHERE serial_number=?",(sn,)).fetchone()
                if exists:
                    conn.execute('UPDATE inventory SET type=?, model=?, tag=?, tombamento=?, condition=?, status=?, assigned_request_id=?, glpi_status=?, last_updated=? WHERE serial_number=?', (w['type'],w['model'],w['tag'],w['tombamento'],w['condition'],w['status'],w['assigned_request_id'],w['glpi_status'],w['last_updated'],sn))
                    upd+=1
                else:
                    conn.execute('INSERT INTO inventory (serial_number,type,model,tag,tombamento,condition,status,assigned_request_id,glpi_status,last_updated) VALUES (?,?,?,?,?,?,?,?,?,?)', (w['serial_number'],w['type'],w['model'],w['tag'],w['tombamento'],w['condition'],w['status'],w['assigned_request_id'],w['glpi_status'],w['last_updated'] or now))
                    ins+=1
        conn.commit()
        out = [[final[s][k] for k in ['serial_number','type','model','tag','tombamento','condition','status','assigned_request_id']] + ["","",final[s]['glpi_status'],final[s].get('last_updated',now)] for s in sorted(final)]
        gsheets.write_sheet(f"{settings.INVENTORY_RANGE}!A2:L", out)
        audit_log(current_user.id, 'Sync', f"I:{ins}, U:{upd}, D:{dele}")
        return jsonify({'success': True, 'message': f"Sync OK! I:{ins}, U:{upd}, D:{dele}"})
    except Exception as e: return jsonify({'message': str(e)}), 500